#!/usr/bin/env python3
"""Append the corrected-architecture pilot addendum to Parker's IPCCC report."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches


HEADING = "Appendix B. Corrected-Architecture Pilot Findings and Expected Impact"


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.space_after = Inches(0.04)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("document", type=Path)
    args = parser.parse_args()

    document = Document(args.document)
    if any(paragraph.text.strip() == HEADING for paragraph in document.paragraphs):
        raise SystemExit(f"Addendum already exists in {args.document}")

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading(HEADING, level=1)
    document.add_paragraph(
        "Status: preliminary pilot evidence only. The corrected full experiment has not yet been run, "
        "so the values below guide the design but do not replace the final inferential results."
    )

    document.add_heading("B.1 Architecture correction and pilot scope", level=2)
    document.add_paragraph(
        "The simulator and allocator boundary were revised because the prior online architecture could "
        "admit pending tasks during completion or allocator calls, destructively reset allocator state, "
        "and expose task information outside the robot messaging path. Those behaviors prevented strict "
        "coalescing bounds from being the treatment actually measured. The corrected architecture uses "
        "exact ordinary batches, a physically gated terminal residual, message-only task discovery, "
        "non-destructive admission, persistent execution bundles, autonomous event-driven recovery, and "
        "allocator-transaction timing that excludes transport and serialization. CBAA remains a "
        "single-current-task allocator but now considers the complete locally known pool and preserves its "
        "winning bid while moving."
    )
    document.add_paragraph(
        "The replacement pilot completed 512 technically valid jobs: 200 rate-confirmation jobs, "
        "288 policy-sweep jobs, and 24 zero-compute diagnostics. It observed no piggyback admissions or "
        "final-release flushes."
    )

    document.add_heading("B.2 Preliminary pilot selections", level=2)
    add_bullet(document, "Final loads: 0.075, 0.30, and 0.60 task arrivals per mission-second.")
    add_bullet(document, "Final policies: Eager, Count B2, Count B4, Count B8, and Bounded B4/W10.")
    add_bullet(
        document,
        "Bounded B4/W10 versus Eager reduced allocator processor work by about 4.9%, 4.3%, and 1.7% "
        "at low, medium, and high load, while increasing mean task latency by about 33.5%, 15.4%, and 5.1%.",
    )
    add_bullet(
        document,
        "Count B4 showed a more aggressive and less uniform tradeoff: work changed by about -10.8%, "
        "+2.8%, and -10.2%, while latency increased by about 135.7%, 26.9%, and 14.4% across the same loads.",
    )
    add_bullet(
        document,
        "The high-load setting remains useful because it exposes liveness sensitivity: B4/W10 completed "
        "12/12 pilot jobs, while Eager completed 11/12. This is a diagnostic signal, not a final ranking.",
    )

    document.add_heading("B.3 Expected changes to the experiment plan", level=2)
    document.add_paragraph(
        "The full AGX design will retain four allocators, three loads, and five policies. Each condition "
        "will receive 50 paired traces in both host-causal and zero-time treatments, executed as 25 traces "
        "for every condition in Round 1 and the remaining 25 traces for every condition in Round 2. "
        "Round 1 verification will occur before Round 2 but will not restrict or select the second-half "
        "matrix. The three-board RP2040 subset remains Eager and Count B4 on four paired traces, for "
        "96 hardware missions total. Old results will not be pooled with the corrected dataset."
    )

    document.add_heading("B.4 Expected changes to the current findings", level=2)
    document.add_paragraph(
        "Numeric conclusions in the current report that depend on the pre-correction reallocation behavior "
        "should be treated as superseded until the full matrix is complete. Under strict coalescing, epoch "
        "counts should become policy-determined rather than completion-driven, latency penalties should be "
        "more visible, and compute savings may remain modest or non-monotonic because larger task pools can "
        "increase the cost of a single allocator transaction. CBAA results may change materially because "
        "movement no longer refreshes its winning bid. Allocator-specific completion rules—suffix release "
        "for ACBBA/HIPC and itemwise repair for PI—may also create legitimate differences in claim churn "
        "and recovery behavior. The corrected experiment is therefore expected to revise both effect sizes "
        "and some algorithm rankings, while more cleanly isolating the actual cost-versus-latency tradeoff "
        "of coalescing."
    )

    document.add_heading("B.5 Final analysis priorities", level=2)
    add_bullet(document, "First establish mission completion, failure classification, and recovery incidence.")
    add_bullet(document, "Use mission-total allocator processor work as the principal compute outcome.")
    add_bullet(document, "Use mean and 95th-percentile release-to-completion latency as the principal responsiveness outcomes.")
    add_bullet(document, "Report mission time and team steps as operational outcomes.")
    add_bullet(
        document,
        "Treat traces—not tasks—as the independent replicates, pair every comparison to Eager on the same "
        "manifest, and do not compare raw bundle-claim counts as if they had identical meaning across allocators.",
    )

    document.save(args.document)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
